"""Импорт staging из Raschet_dlitelnosti.xlsx (P1-B). Не извлекает θ из нормочасов."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app.db.models import EstimateStagingRow
from app.services.cv_adapter.service import canonicalize_class_name
from app.core.config import get_settings
from app.services.knowledge_base.service import load_kb


def _cell(row, idx: int):
    if row is None or idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    s = str(v).strip()
    return s if s and s.lower() != "none" else None


def _guess_equipment_canonical(raw: str | None) -> tuple[str | None, str]:
    if not raw:
        return None, "UNKNOWN"
    text = raw.lower()
    settings = get_settings()
    class_map = {str(k).lower(): str(v) for k, v in (settings.section("detector").get("class_map") or {}).items()}
    # ключевое слово → сырой hint класса
    hints = [
        ("экскаватор", "excavator"),
        ("самосвал", "dump_truck"),
        ("бульдозер", "bulldozer"),
        ("погрузчик", "loader"),
        ("кран", "crane"),
        ("бетононасос", "concrete_pump"),
        ("автобетоносмеситель", "concrete_mixer"),
        ("бетономешалка", "concrete_mixer"),
        ("каток", "roller"),
        ("грейдер", "grader"),
        ("машина", "machinery"),
        ("автомобиль", "truck"),
    ]
    for kw, raw_cls in hints:
        if kw in text:
            canon = canonicalize_class_name(raw_cls, class_map) or raw_cls
            detectable = "NOT_DETECTABLE_BY_CV" if canon.endswith("_unknown") or canon == "machinery" else "DETECTABLE"
            if canon in ("person",):
                detectable = "NOT_DETECTABLE_BY_CV"
            return canon, detectable
    # Фразы только про труд (без техники)
    if any(x in text for x in ("рабочий", "звено", "бригада", "вручную")):
        return None, "NOT_DETECTABLE_BY_CV"
    return None, "UNKNOWN"


def import_estimate_xlsx(db: Session, path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = None
    for name in wb.sheetnames:
        if "смет" in name.lower() or name.strip().endswith("3"):
            sheet_name = name
            break
    if sheet_name is None:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    # Очистить предыдущий staging для этого файла
    db.query(EstimateStagingRow).filter(EstimateStagingRow.source_file == str(path)).delete()

    current_group = None
    inserted = 0
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= 4:
            continue
        # Индексы колонок: C=2, F=5, G=6, L=11, M=12, O=14, P=15, Q=16 (0-based)
        group = _cell(row, 2)
        operation = _cell(row, 5)
        quantity = _cell(row, 6)
        equip_desc = _cell(row, 11)
        equip_count = _cell(row, 12)
        dur_o = _cell(row, 14)
        dur_p = _cell(row, 15)
        dur_q = _cell(row, 16)

        if group:
            current_group = group
        if not operation and not equip_desc:
            continue
        # Пропуск строк только с единицами без текста операции
        if operation and re.fullmatch(r"[\d\s.,]+", operation):
            continue

        equip_raw = " | ".join(x for x in [equip_desc, equip_count] if x)
        canon, detectable = _guess_equipment_canonical(equip_desc or equip_raw)
        db.add(
            EstimateStagingRow(
                source_file=str(path),
                sheet=sheet_name,
                row_no=i,
                group_name=current_group,
                operation=operation,
                equipment_raw=equip_raw or None,
                equipment_canonical=canon,
                quantity=quantity,
                unit=None,
                machine_time=None,
                labor_time=None,
                accepted_duration=dur_q or dur_p or dur_o,
                review_status="pending",
                detectable_by_cv=detectable,
                payload_json=json.dumps(
                    {
                        "duration_variants": [dur_o, dur_p, dur_q],
                        "note": "нормативы часов НЕ используются как θ / P(work)",
                    },
                    ensure_ascii=False,
                ),
            )
        )
        inserted += 1

    db.flush()
    kb = load_kb()
    return {
        "source_file": str(path),
        "sheet": sheet_name,
        "rows_staged": inserted,
        "kb_version": kb.version,
        "note": "staging only; θ не извлекается из машинного времени",
    }
