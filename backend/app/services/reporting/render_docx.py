"""DOCX из общего report DTO (python-docx). Без CV inference. Формулировки по ТЗ отчётов."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Cm, Pt

TZ_DISCLAIMER = (
    "Настоящий отчёт носит информационно-аналитический характер. Он не заменяет предусмотренные "
    "законодательством и контрактом документы о приёмке, экспертизу результатов исполнения контракта, "
    "строительный контроль и решения уполномоченных лиц. Правовые последствия определяются "
    "уполномоченными лицами с учётом условий контракта и применимого законодательства."
)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text or "Не установлено")


def _kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v or "Не установлено"


def _try_picture(doc: Document, fr: dict[str, Any]) -> None:
    from app.services.reporting.render_pdf import _resolve_frame_file

    fid = fr.get("frame_id")
    ts = fr.get("captured_at") or "—"
    src = fr.get("timestamp_source") or "—"
    _p(
        doc,
        f"Кадр #{fid} · время: {ts} · тип времени: {src} · камера: {fr.get('camera_id') or 'Не установлено'}",
    )
    if not fid:
        _p(doc, "Изображение недоступно")
        return
    img = _resolve_frame_file(int(fid))
    if img is None:
        _p(doc, "Изображение недоступно в хранилище отчёта")
        return
    try:
        doc.add_picture(str(img), width=Cm(10))
    except Exception:
        _p(doc, "Не удалось вложить изображение")


def render_finding_docx(dto: dict[str, Any], out_path: Path) -> Path:
    project = dto.get("project") or {}
    finding = dto.get("finding") or {}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    _add_heading(doc, "Отчёт по предупреждению PlanSight")
    _kv_table(
        doc,
        [
            ("Проект", str(project.get("name") or project.get("id"))),
            ("Адрес", str(project.get("address") or "Не установлено")),
            ("Идентификатор предупреждения", str(finding.get("finding_id") or finding.get("deviation_id"))),
            ("Сформирован", str(dto.get("generated_at"))),
            ("Автор", str(dto.get("author") or "PlanSight")),
            ("Код", str(finding.get("code") or "Не установлено")),
            ("Статус рассмотрения", str(finding.get("lifecycle") or "Не установлено")),
            ("Статус связи с КСГ", str(finding.get("match_status") or "Не установлено")),
            ("Решение специалиста", str(finding.get("last_verdict") or "нет")),
        ],
    )

    _add_heading(doc, "Краткий вывод", 2)
    _p(doc, str(finding.get("explanation_ru") or "Не установлено"))

    _add_heading(doc, "Контекст работ и связь с КСГ", 2)
    si = finding.get("schedule_item")
    if si:
        _kv_table(
            doc,
            [
                ("Работа КСГ", str(si.get("name") or "Не установлено")),
                ("Плановое окончание", str(si.get("planned_finish") or si.get("planned_end") or "Не установлено")),
                ("Прогноз", str(si.get("forecast_end") or "Не установлено")),
            ],
        )
    elif finding.get("schedule_item_refused_reason"):
        _p(doc, str(finding["schedule_item_refused_reason"]))
    else:
        _p(doc, "Связь с работой КСГ не подтверждена")

    _add_heading(doc, "Ожидалось и на кадрах", 2)
    expected = ", ".join(str(x) for x in (finding.get("expected") or [])) or "Не установлено"
    observed = ", ".join(str(x) for x in (finding.get("observed") or [])) or "Не установлено"
    _p(doc, f"Ожидается: {expected}")
    _p(doc, f"На кадрах: {observed}")

    _add_heading(doc, "Цепочка проверки", 2)
    tr = finding.get("decision_trace") or {}
    _p(doc, f"Статус: {tr.get('trace_status') or 'Не установлено'}")
    for step in tr.get("user_logic") or []:
        _p(doc, f"{step.get('title')}: {step.get('body')}")

    _add_heading(doc, "Доказательства: кадры", 2)
    frames = finding.get("frames") or []
    if not frames:
        _p(doc, "Нет связанных кадров")
    for fr in frames:
        _try_picture(doc, fr)

    _add_heading(doc, "Результат рассмотрения специалистом", 2)
    hist = finding.get("verdict_history") or []
    if not hist:
        _p(doc, "История решений пуста")
    for v in hist:
        _p(doc, f"{v.get('created_at') or '—'} · {v.get('verdict')} · {v.get('user_id') or '—'}")

    _add_heading(doc, "Рекомендуемые действия", 2)
    _p(doc, str(finding.get("recommendation") or "Подтвердите или отклоните сигнал после проверки."))

    _add_heading(doc, "Ограничения и методика", 2)
    for lim in finding.get("limitations") or ["Не установлено"]:
        _p(doc, f"• {lim}")
    _p(doc, str(dto.get("methodology_note") or ""))

    _add_heading(doc, "Оговорка", 2)
    _p(doc, str(dto.get("disclaimer") or TZ_DISCLAIMER))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def render_project_docx(dto: dict[str, Any], out_path: Path) -> Path:
    project = dto.get("project") or {}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    _add_heading(doc, "Сводный отчёт по проекту PlanSight")
    _kv_table(
        doc,
        [
            ("Проект", str(project.get("name"))),
            (
                "Период выборки",
                f"{(dto.get('period') or {}).get('date_from') or '—'} — {(dto.get('period') or {}).get('date_to') or '—'}",
            ),
            ("Сформирован", str(dto.get("generated_at"))),
            ("Версия КСГ", str(dto.get("schedule_version_id") or "Не установлено")),
        ],
    )

    _add_heading(doc, "Управленческое резюме", 2)
    s = dto.get("summary") or {}
    for k, label in (
        ("total", "Всего предупреждений"),
        ("open", "Открытые / в проверке"),
        ("confirmed_by_specialist", "Подтверждены специалистом"),
        ("rejected_false_positive", "Отклонённые сигналы"),
        ("no_data", "Недостаточно данных"),
        ("cv_verified", "С кадрами наблюдения"),
    ):
        _p(doc, f"{label}: {s.get(k, 0)}")

    _add_heading(doc, "Реестр предупреждений", 2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, name in enumerate(["ID", "Код", "Статус", "Связь с КСГ", "Кадры", "Кратко"]):
        hdr[i].text = name
    for row in dto.get("registry") or []:
        cells = table.add_row().cells
        cells[0].text = str(row.get("finding_id"))
        cells[1].text = str(row.get("code") or "")
        cells[2].text = str(row.get("lifecycle") or "")
        cells[3].text = str(row.get("match_status") or "")
        cells[4].text = str(row.get("evidence_frames") or 0)
        cells[5].text = str(row.get("explanation_ru") or "")[:180]

    for f in dto.get("findings") or []:
        _add_heading(doc, f"Предупреждение #{f.get('finding_id')}", 2)
        _p(doc, str(f.get("explanation_ru") or "Не установлено"))
        for fr in (f.get("frames") or [])[:3]:
            _try_picture(doc, fr)

    _add_heading(doc, "Методика и ограничения", 2)
    _p(doc, str(dto.get("methodology_note") or ""))
    _p(doc, f"Фильтры: {dto.get('filters')}")

    _add_heading(doc, "Оговорка", 2)
    _p(doc, str(dto.get("disclaimer") or TZ_DISCLAIMER))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
